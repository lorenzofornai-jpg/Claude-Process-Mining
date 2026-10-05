"""Contenuto dei documenti di supporto come contesto per il mapping AI.

Data dictionary, manuali e procedure del cliente spiegano cosa significano
tabelle custom, campi e codici (es. ZZ_AR_RISK_LOG.REASON_CODE): proprio
cio' che Claude non puo' dedurre dai soli nomi e da pochi valori di esempio.

1. All'upload (o la prima volta che serve, per i documenti gia' caricati) si
   estrae il testo: PDF, Word (.docx), PowerPoint (.pptx), Excel (.xlsx),
   CSV/TXT/Markdown, XML/BPMN. Immagini, disegni (.vsdx, .drawio, .svg) e i
   vecchi formati Office (.doc, .xls) non hanno testo leggibile qui.
2. Prima del mapping si scelgono, senza AI e senza costo, i brani che citano
   la tabella, le sue colonne o i codici delle sue colonne attivita', entro
   un tetto di caratteri per tabella: a Claude arrivano solo quelli, non i
   documenti interi (costi prevedibili, contesto mirato).
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_TEXT_CHARS = 400_000          # testo conservato per documento
CHUNK_CHARS = 700                 # dimensione di un brano
EXCERPT_CHARS_PER_TABLE = 2_500   # brani per tabella nella chiamata di mapping
SKELETON_CHARS_PER_TABLE = 600    # brani per tabella nella chiamata sul modello comune
TEXT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md", ".xml", ".bpmn"}
DOC_TYPE_WEIGHT = {"data_dictionary": 2.0, "manual": 1.3, "procedure": 1.2, "bpmn": 1.0, "other": 1.0}


# ---------------------------------------------------------------- estrazione
def _docx(path: Path) -> str:
    """Paragrafi e tabelle (una riga di tabella = celle separate da " | ")."""
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    lines: list[str] = []

    def text(el) -> str:
        return "".join(t.text or "" for t in el.iter(w + "t"))

    body = root.find(w + "body")
    for el in (body if body is not None else []):
        if el.tag == w + "p":
            lines.append(text(el))
        elif el.tag == w + "tbl":
            for tr in el.iter(w + "tr"):
                lines.append(" | ".join(text(tc).strip() for tc in tr.findall(w + "tc")))
    return "\n".join(lines)


def _pptx(path: Path) -> str:
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    out = []
    with zipfile.ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        for i, name in enumerate(slides, start=1):
            root = ET.fromstring(z.read(name))
            paras = ["".join(t.text or "" for t in p.iter(a + "t")) for p in root.iter(a + "p")]
            out.append(f"[Slide {i}]\n" + "\n".join(p for p in paras if p.strip()))
    return "\n\n".join(out)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
            if sum(len(r) for r in rows) > MAX_TEXT_CHARS:
                break
        if rows:
            out.append(f"[Foglio {ws.title}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(out)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        pages.append(f"[Pagina {i}]\n{page.extract_text() or ''}")
        if sum(len(p) for p in pages) > MAX_TEXT_CHARS:
            break
    return "\n\n".join(pages)


def _xml(path: Path) -> str:
    root = ET.parse(path).getroot()
    parts = []
    for el in root.iter():
        for attr in ("name", "label"):
            if el.get(attr):
                parts.append(el.get(attr))
        if el.text and el.text.strip():
            parts.append(el.text.strip())
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    """Testo leggibile del documento ("" se il formato non ha testo o la lettura fallisce)."""
    ext = path.suffix.lower()
    if ext not in TEXT_EXTENSIONS:
        return ""
    try:
        if ext == ".pdf":
            text = _pdf(path)
        elif ext == ".docx":
            text = _docx(path)
        elif ext == ".pptx":
            text = _pptx(path)
        elif ext == ".xlsx":
            text = _xlsx(path)
        elif ext in (".xml", ".bpmn"):
            text = _xml(path)
        else:
            text = path.read_bytes().decode("utf-8", errors="replace")
    except Exception as exc:  # un documento illeggibile non deve mai bloccare l'assessment
        print(f"Testo non estraibile da {path.name} ({exc!r}).")
        return ""
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()
    return text[:MAX_TEXT_CHARS]


# ---------------------------------------------------------------- brani
def _chunks(text: str) -> list[str]:
    """Brani di circa CHUNK_CHARS caratteri che non spezzano le righe (una riga di
    data dictionary, cioe' campo + descrizione, resta intera)."""
    out, cur = [], ""
    for line in text.splitlines():
        if len(cur) + len(line) + 1 > CHUNK_CHARS and cur:
            out.append(cur.strip())
            cur = ""
        cur += line[: CHUNK_CHARS * 2] + "\n"
    if cur.strip():
        out.append(cur.strip())
    return out


def _term_re(term: str) -> re.Pattern:
    return re.compile(r"(?<![A-Za-z0-9_])" + re.escape(term) + r"(?![A-Za-z0-9_])", re.I)


def _terms_for(table) -> tuple[list[str], list[str]]:
    """(termini forti, termini deboli) di una tabella: il nome tabella e i codici
    delle colonne attivita' sono forti; i nomi colonna sono deboli (spesso generici)."""
    strong = [table["name"]]
    strong += [v for v in table.get("activity_values", []) if len(v) >= 4 and not v.isdigit()]
    weak = [c for c in table["columns"] if len(c) >= 3]
    return strong, weak


def select_excerpts(docs: list[dict], tables: list[dict]) -> dict:
    """Brani pertinenti per tabella, scelti senza AI.

    docs: [{"filename", "doc_type", "text"}]; tables: [{"name", "columns": [...],
    "activity_values": [...]}]. Ritorna {"by_table": {tabella: [{"doc", "text"}]},
    "used_docs": [...], "excerpt_count": n}. Un brano conta se cita il nome della
    tabella o un suo codice, oppure almeno due sue colonne (una colonna da sola,
    es. BELNR, compare in troppe tabelle per essere un indizio)."""
    chunks = [(d["filename"], DOC_TYPE_WEIGHT.get(d.get("doc_type"), 1.0), c)
              for d in docs if d.get("text") for c in _chunks(d["text"])]
    by_table: dict[str, list[dict]] = {}
    used: set[str] = set()
    for t in tables:
        strong, weak = _terms_for(t)
        strong_re = [_term_re(s) for s in strong]
        weak_re = [_term_re(w) for w in weak]
        scored = []
        for i, (doc, weight, text) in enumerate(chunks):
            hits_strong = sum(1 for r in strong_re if r.search(text))
            hits_weak = sum(1 for r in weak_re if r.search(text))
            if hits_strong == 0 and hits_weak < 2:
                continue
            scored.append(((3 * hits_strong + hits_weak) * weight, i, doc, text))
        picked, size = [], 0
        for score, i, doc, text in sorted(scored, key=lambda x: (-x[0], x[1])):
            if size + len(text) > EXCERPT_CHARS_PER_TABLE:
                continue
            picked.append((i, doc, text))
            size += len(text)
        if picked:
            # nell'ordine del documento: brani consecutivi si leggono come un testo unico
            by_table[t["name"]] = [{"doc": doc, "text": text} for i, doc, text in sorted(picked)]
            used.update(doc for _, doc, _ in picked)
    return {"by_table": by_table, "used_docs": sorted(used),
            "excerpt_count": sum(len(v) for v in by_table.values())}


def for_skeleton(by_table: dict[str, list[dict]]) -> dict[str, list[str]]:
    """Versione corta dei brani per la chiamata sul modello comune."""
    out = {}
    for table, items in by_table.items():
        texts, size = [], 0
        for it in items:
            if size >= SKELETON_CHARS_PER_TABLE:
                break
            piece = it["text"][: SKELETON_CHARS_PER_TABLE - size]
            texts.append(piece)
            size += len(piece)
        out[table] = texts
    return out
